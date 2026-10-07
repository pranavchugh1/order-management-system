import asyncio
import json
import os
import tempfile
import anyio

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from starlette.background import BackgroundTask
from starlette.concurrency import run_in_threadpool

from auth import current_user, require_page
from database import db, now_iso
from date_format import display_date

router = APIRouter(prefix='/api/exports', dependencies=[Depends(current_user)])
export_slots = asyncio.Semaphore(2)
HEADER_FONT = Font(bold=True, color='FFFFFF')
HEADER_FILL = PatternFill('solid', fgColor='6F451F')
PARTY_FONT = Font(bold=True, color='1F1912')
PARTY_FILL = PatternFill('solid', fgColor='F5E6D0')


def append_row(sheet, values, style=None):
    cells = []
    for value in values:
        cell = WriteOnlyCell(sheet, value=value)
        if isinstance(value, str):
            cell.data_type = 's'
        if style == 'header':
            cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        elif style == 'party':
            cell.font, cell.fill = PARTY_FONT, PARTY_FILL
        cells.append(cell)
    sheet.append(cells)


def new_sheet(book, name, headers, widths=None):
    sheet = book.create_sheet(name)
    sheet.freeze_panes = 'A2'
    for i, title in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(i)].width = (widths or {}).get(title, 20)
    append_row(sheet, headers, 'header')
    return sheet


def finalize(sheet, rows, cols):
    sheet.auto_filter.ref = f'A1:{get_column_letter(cols)}{max(1, rows + 1)}'


def parcel_contents(parcel):
    return ' + '.join(f"{c['catalogue_name']} {c['volume_name']} ({c['sets']} sets)" for c in parcel['compositions'])


async def stream_orders(state):
    """Stream visible orders for one state, ordered by party then created_at."""
    projection = {'_id': 0, 'id': 1, 'order_number': 1, 'date': 1, 'created_at': 1, 'party_name': 1, 'party_id': 1, 'parcels': 1, 'remarks': 1}
    cursor = db.orders.find({'visible': True, 'state': state}, projection).sort([('party_name_key', 1), ('created_at', 1), ('id', 1)]).batch_size(200)
    async for order in cursor:
        yield order


def build_pending_requirements(snapshot, filename, generated):
    book = Workbook(write_only=True)
    summary = new_sheet(book, 'Summary', ['Metric', 'Value'], {'Metric': 30, 'Value': 30})
    volumes = new_sheet(book, 'Volume Requirements', ['Catalogue', 'Volume', 'Pieces per Set', 'Pending Sets', 'Pending Pieces'], {'Catalogue': 34})
    catalogues = new_sheet(book, 'Catalogue Totals', ['Catalogue', 'Pending Sets', 'Pending Pieces'], {'Catalogue': 34})
    volume_map, catalogue_map = {}, {}
    total_sets = total_pieces = orders_count = parcels_count = 0
    with open(snapshot, encoding='utf-8') as source:
        for line in source:
            order = json.loads(line)
            orders_count += 1
            for parcel in order['parcels']:
                if parcel['status'] == 'Parcel Sent':
                    continue
                parcels_count += 1
                for c in parcel['compositions']:
                    pieces = c['sets'] * c['pieces_per_set']
                    v = volume_map.setdefault((c['catalogue_id'], c['volume_id']), [c['catalogue_name'], c['volume_name'], c['pieces_per_set'], 0, 0])
                    v[3] += c['sets']; v[4] += pieces
                    cat = catalogue_map.setdefault(c['catalogue_id'], [c['catalogue_name'], 0, 0])
                    cat[1] += c['sets']; cat[2] += pieces
                    total_sets += c['sets']; total_pieces += pieces
    for row in sorted(volume_map.values(), key=lambda x: (x[0].casefold(), x[1].casefold())):
        append_row(volumes, row)
    for row in sorted(catalogue_map.values(), key=lambda x: x[0].casefold()):
        append_row(catalogues, row)
    for row in [('Generated at (UTC)', display_date(generated, True)), ('Pending orders', orders_count), ('Pending parcels', parcels_count), ('Distinct catalogue+volume rows', len(volume_map)), ('Pending sets', total_sets), ('Pending pieces', total_pieces)]:
        append_row(summary, row)
    finalize(volumes, len(volume_map), 5)
    finalize(catalogues, len(catalogue_map), 3)
    book.save(filename)


def build_orders_workbook(snapshot, filename, generated, mode):
    book = Workbook(write_only=True)
    summary = new_sheet(book, 'Summary', ['Metric', 'Value'], {'Metric': 32, 'Value': 32})
    if mode == 'pending':
        headers = ['Party', 'Order', 'Order Date', 'Created At', 'Parcel Pieces', 'Contents', 'Parcel Status', 'Remarks']
    else:
        headers = ['Party', 'Order', 'Order Date', 'Parcel Pieces', 'Contents', 'Challan #', 'Firm', 'Sent Date', 'Remarks']
    widths = {'Party': 34, 'Contents': 60, 'Firm': 32, 'Remarks': 34, 'Created At': 22, 'Sent Date': 22}
    sheet = new_sheet(book, 'Pending Orders' if mode == 'pending' else 'Completed Orders', headers, widths)
    row_count = 0; orders_count = parcels_count = pieces_total = 0; last_party_id = None; parties = set()
    with open(snapshot, encoding='utf-8') as source:
        for line in source:
            order = json.loads(line)
            parties.add(order['party_id']); orders_count += 1
            if order['party_id'] != last_party_id:
                banner = [f"Party - {order['party_name']}"] + [''] * (len(headers) - 1)
                append_row(sheet, banner, 'party'); row_count += 1
                last_party_id = order['party_id']
            for parcel in order['parcels']:
                if mode == 'pending' and parcel['status'] == 'Parcel Sent':
                    continue
                if mode == 'completed' and parcel['status'] != 'Parcel Sent':
                    continue
                parcels_count += 1
                pieces_total += sum(c['sets'] * c['pieces_per_set'] for c in parcel['compositions'])
                contents = parcel_contents(parcel)
                if mode == 'pending':
                    append_row(sheet, [order['party_name'], order['order_number'], display_date(order['date']), display_date(order.get('created_at'), True), parcel['target'], contents, parcel['status'], order.get('remarks', '')])
                else:
                    append_row(sheet, [order['party_name'], order['order_number'], display_date(order['date']), parcel['target'], contents, parcel.get('challan_number', ''), parcel.get('firm', ''), display_date(parcel.get('sent_at'), True), order.get('remarks', '')])
                row_count += 1
    for row in [('Generated at (UTC)', display_date(generated, True)), ('Scope', 'Pending orders and parcels' if mode == 'pending' else 'Fully dispatched orders'), ('Parties', len(parties)), ('Orders', orders_count), ('Parcels', parcels_count), ('Total pieces', pieces_total)]:
        append_row(summary, row)
    finalize(sheet, row_count, len(headers))
    book.save(filename)


async def snapshot_orders(state, snapshot):
    rows_total = 0
    buffer = []
    async with await anyio.open_file(snapshot, 'w', encoding='utf-8') as target:
        async for order in stream_orders(state):
            rows_total += sum(len(p['compositions']) for p in order['parcels'])
            if rows_total > 1000000:
                raise HTTPException(413, 'The report exceeds the Excel row limit. Contact the administrator to split the report.')
            buffer.append(json.dumps(order) + '\n')
            if len(buffer) >= 100:
                await target.write(''.join(buffer))
                buffer.clear()
        if buffer:
            await target.write(''.join(buffer))


async def guarded(builder, filename_prefix):
    if export_slots.locked():
        raise HTTPException(429, 'Other reports are being generated. Please try again shortly.', headers={'Retry-After': '10'})
    async with export_slots:
        directory = tempfile.TemporaryDirectory(prefix='aditya-export-')
        try:
            path = await builder(directory.name)
            return FileResponse(path, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', filename=f'{filename_prefix}-{display_date(now_iso())}.xlsx', background=BackgroundTask(directory.cleanup))
        except Exception:
            directory.cleanup()
            raise


@router.get('/pending-requirements.xlsx')
async def export_pending_requirements(user=Depends(require_page('pending_requirements'))):
    async def build(tmp):
        snapshot = os.path.join(tmp, 'orders.jsonl')
        filename = os.path.join(tmp, 'pending-requirements.xlsx')
        await snapshot_orders('pending', snapshot)
        await run_in_threadpool(build_pending_requirements, snapshot, filename, now_iso())
        return filename
    return await guarded(build, 'Aditya-Prints-Pending-Requirements')


@router.get('/pending-orders.xlsx')
async def export_pending_orders(user=Depends(require_page('orders_pending'))):
    async def build(tmp):
        snapshot = os.path.join(tmp, 'orders.jsonl')
        filename = os.path.join(tmp, 'pending-orders.xlsx')
        await snapshot_orders('pending', snapshot)
        await run_in_threadpool(build_orders_workbook, snapshot, filename, now_iso(), 'pending')
        return filename
    return await guarded(build, 'Aditya-Prints-Pending-Orders')


@router.get('/completed-orders.xlsx')
async def export_completed_orders(user=Depends(require_page('orders_completed'))):
    async def build(tmp):
        snapshot = os.path.join(tmp, 'orders.jsonl')
        filename = os.path.join(tmp, 'completed-orders.xlsx')
        await snapshot_orders('completed', snapshot)
        await run_in_threadpool(build_orders_workbook, snapshot, filename, now_iso(), 'completed')
        return filename
    return await guarded(build, 'Aditya-Prints-Completed-Orders')


# Backwards-compatible alias kept for any bookmarked links; delegates to the pending orders sheet.
@router.get('/pending.xlsx')
async def export_pending_alias(user=Depends(require_page('orders_pending'))):
    return await export_pending_orders()
