"""Iteration 7 targeted regression: fixture setup for UI flow + XLSX/date validations."""

import io
import json
import os
import re
import uuid

import pytest
import requests
from dotenv import load_dotenv
from openpyxl import load_workbook


load_dotenv("/app/frontend/.env")
load_dotenv("/app/backend/.env")

BASE_URL = os.environ["REACT_APP_BACKEND_URL"].rstrip("/")
ADMIN_EMAIL = os.environ["ADMIN_EMAIL"]
ADMIN_PASSWORD = os.environ["ADMIN_PASSWORD"]
STATE_PATH = "/app/test_reports/iteration7_state.json"
FIRMS = ("Dhan Guru Nanak Synthetics", "Alveera Fashion Private Limited")

DATE_RE = re.compile(r"^\d{2}-\d{2}-\d{4}$")
DATETIME_RE = re.compile(r"^\d{2}-\d{2}-\d{4}(?:\s+\d{2}:\d{2}:\d{2})?$")


def _read_state():
    if not os.path.exists(STATE_PATH):
        return {}
    with open(STATE_PATH, "r", encoding="utf-8") as fh:
        return json.load(fh)


def _write_state(data):
    with open(STATE_PATH, "w", encoding="utf-8") as fh:
        json.dump(data, fh, indent=2)


def _upsert_state(change):
    current = _read_state()
    current.update(change)
    _write_state(current)


@pytest.fixture(scope="module")
def api_and_seed():
    """Auth + isolated fixture creation for requested 10/6 catalogue flow."""
    s = requests.Session()
    s.headers.update({"Content-Type": "application/json", "Origin": BASE_URL, "Referer": BASE_URL + "/"})
    login = s.post(f"{BASE_URL}/api/auth/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD})
    assert login.status_code == 200, f"login failed: {login.status_code} {login.text}"
    csrf = s.cookies.get("csrf_token")
    assert csrf, "csrf cookie missing"
    s.headers.update({"X-CSRF-Token": csrf})

    suffix = uuid.uuid4().hex[:8]
    party = s.post(f"{BASE_URL}/api/parties", json={"name": f"TEST Party {suffix}"})
    assert party.status_code == 200, party.text
    party = party.json()

    cat = s.post(f"{BASE_URL}/api/catalogues", json={"name": f"TEST Cat {suffix}"})
    assert cat.status_code == 200, cat.text
    cat = cat.json()

    vol_six = s.post(
        f"{BASE_URL}/api/catalogues/{cat['id']}/volumes",
        json={"name": "Six", "pieces_per_set": 6},
    )
    assert vol_six.status_code == 200, vol_six.text
    vol_six = vol_six.json()

    vol_ten = s.post(
        f"{BASE_URL}/api/catalogues/{cat['id']}/volumes",
        json={"name": "Ten", "pieces_per_set": 10},
    )
    assert vol_ten.status_code == 200, vol_ten.text
    vol_ten = vol_ten.json()

    # Additional distinct 6-pcs volume for the "two six volumes @72" frontend guardrail check.
    vol_six_alt = s.post(
        f"{BASE_URL}/api/catalogues/{cat['id']}/volumes",
        json={"name": "Six Alt", "pieces_per_set": 6},
    )
    assert vol_six_alt.status_code == 200, vol_six_alt.text
    vol_six_alt = vol_six_alt.json()

    # Build one pending and one completed order (date pinned for DD-MM-YYYY checks).
    pending_payload = {
        "party_id": party["id"],
        "party_name": party["name"],
        "date": "2026-02-20",
        "remarks": "TEST pending iteration7",
        "parcels": [
            {
                "target": 72,
                "status": "Pending",
                "compositions": [
                    {
                        "catalogue_id": cat["id"],
                        "catalogue_name": cat["name"],
                        "volume_id": vol_six["id"],
                        "volume_name": vol_six["name"],
                        "pieces_per_set": 6,
                        "sets": 12,
                    }
                ],
            }
        ],
    }
    pending = s.post(f"{BASE_URL}/api/orders", json=pending_payload)
    assert pending.status_code == 201, pending.text
    pending = pending.json()

    completed_payload = {
        "party_id": party["id"],
        "party_name": party["name"],
        "date": "2026-02-20",
        "remarks": "TEST completed iteration7",
        "parcels": [
            {
                "target": 76,
                "status": "Pending",
                "compositions": [
                    {
                        "catalogue_id": cat["id"],
                        "catalogue_name": cat["name"],
                        "volume_id": vol_ten["id"],
                        "volume_name": vol_ten["name"],
                        "pieces_per_set": 10,
                        "sets": 4,
                    },
                    {
                        "catalogue_id": cat["id"],
                        "catalogue_name": cat["name"],
                        "volume_id": vol_six["id"],
                        "volume_name": vol_six["name"],
                        "pieces_per_set": 6,
                        "sets": 6,
                    },
                ],
            }
        ],
    }
    completed = s.post(f"{BASE_URL}/api/orders", json=completed_payload)
    assert completed.status_code == 201, completed.text
    completed = completed.json()

    sent = s.patch(
        f"{BASE_URL}/api/orders/{completed['id']}/parcels/0",
        json={"version": completed["version"], "challan_number": f"TEST-COMP-{suffix}", "firm": FIRMS[1]},
    )
    assert sent.status_code == 200, sent.text

    # Dedicated pending order with 2 parcels for UI dispatch memory flow.
    ui_dispatch_payload = {
        "party_id": party["id"],
        "party_name": party["name"],
        "date": "2026-02-20",
        "remarks": "TEST ui dispatch memory",
        "parcels": [
            {
                "target": 72,
                "status": "Pending",
                "compositions": [
                    {
                        "catalogue_id": cat["id"],
                        "catalogue_name": cat["name"],
                        "volume_id": vol_six["id"],
                        "volume_name": vol_six["name"],
                        "pieces_per_set": 6,
                        "sets": 12,
                    }
                ],
            },
            {
                "target": 72,
                "status": "Pending",
                "compositions": [
                    {
                        "catalogue_id": cat["id"],
                        "catalogue_name": cat["name"],
                        "volume_id": vol_six["id"],
                        "volume_name": vol_six["name"],
                        "pieces_per_set": 6,
                        "sets": 12,
                    }
                ],
            },
        ],
    }
    ui_dispatch = s.post(f"{BASE_URL}/api/orders", json=ui_dispatch_payload)
    assert ui_dispatch.status_code == 201, ui_dispatch.text
    ui_dispatch = ui_dispatch.json()

    _upsert_state(
        {
            "run_tag": suffix,
            "party": party,
            "catalogue": {"id": cat["id"], "name": cat["name"], "volumes": {"six": vol_six, "ten": vol_ten}},
            "catalogue_extra": {"volumes": {"six_alt": vol_six_alt}},
            "seed_order_ids": [pending["id"], completed["id"], ui_dispatch["id"]],
            "seed_order_numbers": [pending["order_number"], completed["order_number"], ui_dispatch["order_number"]],
            "ui_dispatch_order_id": ui_dispatch["id"],
            "ui_dispatch_order_number": ui_dispatch["order_number"],
            "base_url": BASE_URL,
            "cleanup": {
                "party_ids": [party["id"]],
                "catalogue_ids": [cat["id"]],
                "order_ids": [pending["id"], completed["id"], ui_dispatch["id"]],
            },
        }
    )

    return {
        "session": s,
        "party": party,
        "catalogue": cat,
        "vol_six": vol_six,
        "vol_ten": vol_ten,
        "vol_six_alt": vol_six_alt,
        "pending": pending,
        "completed": completed,
    }


def test_mixed_preview_10_plus_6_for_76(api_and_seed):
    """Bulk parser: quoted references with same catalogue + different volume allocates 4+6 for 76."""
    s = api_and_seed["session"]
    party = api_and_seed["party"]
    cat = api_and_seed["catalogue"]
    six = api_and_seed["vol_six"]
    ten = api_and_seed["vol_ten"]
    text = f'"{cat["name"]} | {ten["name"]}" + "{cat["name"]} | {six["name"]}" 1 parcel @ 76 pcs'
    r = s.post(f"{BASE_URL}/api/bulk/preview", json={"blocks": [{"party_id": party["id"], "text": text, "remarks": "TEST i7"}]})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["errors"] == []
    assert len(body["orders"]) == 1 and len(body["orders"][0]["parcels"]) == 1
    comps = body["orders"][0]["parcels"][0]["compositions"]
    assert [c["sets"] for c in comps] == [4, 6]
    assert sum(c["sets"] * c["pieces_per_set"] for c in comps) == 76


@pytest.mark.parametrize(
    "endpoint,sheet_name,required_date_headers",
    [
        ("/api/exports/pending-requirements.xlsx", "Summary", []),
        ("/api/exports/pending-orders.xlsx", "Pending Orders", ["Order Date", "Created At"]),
        ("/api/exports/completed-orders.xlsx", "Completed Orders", ["Order Date", "Sent Date"]),
    ],
)
def test_exports_filename_headers_and_dates(api_and_seed, endpoint, sheet_name, required_date_headers):
    """XLSX exports: no Parcel # column and DD-MM-YYYY date formatting for key fields."""
    s = api_and_seed["session"]
    r = s.get(f"{BASE_URL}{endpoint}")
    assert r.status_code == 200, r.text
    disposition = r.headers.get("content-disposition", "")
    assert re.search(r"Aditya-Prints-.*-\d{2}-\d{2}-\d{4}\.xlsx", disposition)

    wb = load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
    ws = wb[sheet_name]

    rows = list(ws.iter_rows(values_only=True))
    assert rows, "workbook sheet is empty"

    if sheet_name in {"Pending Orders", "Completed Orders"}:
        headers = [str(h or "").strip() for h in rows[0]]
        assert "Parcel #" not in headers
        for header in required_date_headers:
            assert header in headers
        header_idx = {h: i for i, h in enumerate(headers)}

        state = _read_state()
        expected_completed_order = (state.get("seed_order_numbers") or [None, None])[1]
        data_rows = [r for r in rows[1:] if r and any(r)]
        # ignore party banner rows and blank rows
        order_rows = [r for r in data_rows if isinstance(r[header_idx["Order"]], str) and r[header_idx["Order"]].startswith("ORD-")]
        assert order_rows, "no order rows found"

        if sheet_name == "Completed Orders" and expected_completed_order:
            order_rows = [r for r in order_rows if r[header_idx["Order"]] == expected_completed_order]
            assert order_rows, f"seed completed order missing from export: {expected_completed_order}"

        for header in required_date_headers:
            value = str(order_rows[0][header_idx[header]] or "")
            assert DATETIME_RE.match(value), f"invalid {header}: {value}"
            assert DATE_RE.match(value[:10]), f"date prefix invalid for {header}: {value}"
    else:
        # Summary sheet from pending-requirements.xlsx
        key_values = {str(r[0] or "").strip(): str(r[1] or "").strip() for r in rows if r and len(r) > 1}
        generated = key_values.get("Generated at (UTC)", "")
        assert DATETIME_RE.match(generated), f"generated timestamp invalid: {generated}"
        assert DATE_RE.match(generated[:10])
